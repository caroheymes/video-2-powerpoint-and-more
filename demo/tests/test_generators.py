import pytest
from pathlib import Path
from docx import Document
from pptx import Presentation
import pipeline


class TestDocumentGenerators:
    """Tests des moteurs de génération de documents (Word, PowerPoint, HTML, PDF)."""

    def test_generate_docx_validity(self, tmp_path, mock_llm_data, keyframes_dir):
        """Vérifie la création d'un document Word .docx structuré et lisible."""
        out_docx = tmp_path / "test_article.docx"
        pipeline.generate_docx_guide(mock_llm_data, keyframes_dir, out_docx)

        assert out_docx.exists()
        assert out_docx.stat().st_size > 1000

        doc = Document(str(out_docx))
        full_text = "\n".join([p.text for p in doc.paragraphs])
        assert "Gestion du Temps et Créativité en Entreprise" in full_text
        assert "Le défi de la surcharge d'information" in full_text
        assert len(doc.tables) >= 1

    def test_generate_pptx_validity(self, tmp_path, mock_llm_data, keyframes_dir):
        """Vérifie la génération d'un diaporama PowerPoint 16:9 aux normes."""
        out_pptx = tmp_path / "test_presentation.pptx"
        pipeline.generate_pptx(mock_llm_data, keyframes_dir, out_pptx)

        assert out_pptx.exists()
        assert out_pptx.stat().st_size > 5000

        prs = Presentation(str(out_pptx))
        assert abs(prs.slide_width - 12192000) < 1000
        assert abs(prs.slide_height - 6858000) < 1000
        assert len(prs.slides) == len(mock_llm_data["slides"])

    def test_generate_html_deck_validity(self, tmp_path, mock_llm_data, keyframes_dir):
        """Vérifie la création d'un deck web HTML autonome avec images embarquées."""
        out_html = tmp_path / "test_presentation.html"
        pipeline.generate_html_deck(mock_llm_data, keyframes_dir, out_html)

        assert out_html.exists()
        html_content = out_html.read_text(encoding="utf-8")

        assert "<!DOCTYPE html>" in html_content
        assert "Retrouver du Temps pour Créer" in html_content
        assert "data:image/jpeg;base64," in html_content

    def test_export_pdf_graceful(self, tmp_path, mock_llm_data, keyframes_dir):
        """Vérifie que l'export PDF s'exécute ou échoue gracieusement si aucun Chromium headless n'est installé."""
        out_html = tmp_path / "test_presentation.html"
        out_pdf = tmp_path / "test_presentation.pdf"
        pipeline.generate_html_deck(mock_llm_data, keyframes_dir, out_html)

        try:
            pipeline.export_pdf_headless(out_html, out_pdf)
            if out_pdf.exists():
                assert out_pdf.stat().st_size > 1000
        except Exception as e:
            assert "Chromium" in str(e) or "navigateur" in str(e) or "Edge" in str(e) or "introuvable" in str(e)
